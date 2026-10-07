# external libraries imports
import csv
import io
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

# internal application code imports
from employees.enums import EMPLOYEE_FIELD_GROUPS, EmployeeFieldGroup

from .EmployeeExportColumn import EmployeeExportColumn


# main code
COLUMNS = [
    EmployeeExportColumn('status', 'Estado', 'choice'),
    EmployeeExportColumn('id_type', 'Tipo de identificación', 'choice'),
    EmployeeExportColumn('id_number', 'Identificación', 'integer'),
    EmployeeExportColumn('full_name', 'Apellidos y nombres'),
    EmployeeExportColumn('birth_date', 'Fecha de nacimiento', 'date'),
    EmployeeExportColumn('age', 'Edad', 'integer'),
    EmployeeExportColumn('sex', 'Sexo', 'choice'),
    EmployeeExportColumn('blood_type', 'Grupo sanguíneo', 'choice'),
    EmployeeExportColumn('marital_status', 'Estado civil', 'choice'),
    EmployeeExportColumn('has_children', 'Tiene hijos', 'boolean'),
    EmployeeExportColumn('mobile_phone', 'Celular'),
    EmployeeExportColumn('personal_email', 'Correo electrónico'),
    EmployeeExportColumn('address', 'Dirección de residencia'),
    EmployeeExportColumn('neighborhood', 'Barrio'),
    EmployeeExportColumn('city', 'Ciudad'),
    EmployeeExportColumn('education_level', 'Nivel educativo', 'choice'),
    EmployeeExportColumn('degree_title', 'Título'),
    EmployeeExportColumn('employment_type', 'Tipo de vinculación', 'choice'),
    EmployeeExportColumn('category', 'Categoría', 'choice'),
    EmployeeExportColumn('division_name', 'Dirección', 'relation', 'division.name'),
    EmployeeExportColumn('evaluation_group', 'Grupo de evaluación', 'choice'),
    EmployeeExportColumn('collective_agreement', 'Pacto colectivo', 'boolean'),
    EmployeeExportColumn('position_name', 'Cargo actual', 'relation', 'position.name'),
    EmployeeExportColumn('position_start_date', 'Inicio del cargo actual', 'date'),
    EmployeeExportColumn(
        'previous_position_name', 'Cargo anterior', 'relation', 'previous_position.name'
    ),
    EmployeeExportColumn('previous_position_start_date', 'Inicio del cargo anterior', 'date'),
    EmployeeExportColumn('previous_position_end_date', 'Fin del cargo anterior', 'date'),
    EmployeeExportColumn('is_leader', 'Es líder', 'boolean'),
    EmployeeExportColumn('section_name', 'Sección', 'relation', 'section.name'),
    EmployeeExportColumn('cost_center_code', 'Centro de costos', 'relation', 'cost_center.code'),
    EmployeeExportColumn(
        'cost_center_name', 'Nombre centro de costos', 'relation', 'cost_center.name'
    ),
    EmployeeExportColumn('area', 'Área', 'choice'),
    EmployeeExportColumn('additional_role', 'Rol adicional', 'choice'),
    EmployeeExportColumn(
        'immediate_boss_name', 'Jefe inmediato', 'relation', 'immediate_boss.full_name'
    ),
    EmployeeExportColumn('hire_date', 'Fecha de ingreso', 'date'),
    EmployeeExportColumn('retirement_date', 'Fecha de retiro', 'date'),
    EmployeeExportColumn('seniority', 'Antigüedad', 'seniority'),
    EmployeeExportColumn('training', 'Formación'),
    EmployeeExportColumn('current_salary', 'Salario actual', 'money'),
    EmployeeExportColumn('salary_type', 'Tipo de salario', 'choice'),
    EmployeeExportColumn('hourly_rate', 'Valor hora', 'money'),
    EmployeeExportColumn('transport_allowance', 'Auxilio de transporte', 'money'),
    EmployeeExportColumn('contract_type', 'Tipo de contrato', 'choice'),
    EmployeeExportColumn('contract_end_date', 'Fecha de vencimiento', 'date'),
    EmployeeExportColumn('extensions', 'Prórrogas', 'extensions'),
    EmployeeExportColumn('indefinite_extension', 'Prórroga indefinido'),
    EmployeeExportColumn('occupational_risk_insurer', 'ARL', 'choice'),
    EmployeeExportColumn('health_insurer', 'EPS', 'choice'),
    EmployeeExportColumn('pension_fund', 'Fondo de pensión', 'choice'),
    EmployeeExportColumn('severance_fund', 'Fondo de cesantías', 'choice'),
    EmployeeExportColumn('notes', 'Alertas / Observaciones'),
    EmployeeExportColumn('birth_municipality', 'Municipio de nacimiento'),
    EmployeeExportColumn('nationality', 'Nacionalidad'),
    EmployeeExportColumn('ethnicity', 'Pertenencia étnica', 'choice'),
    EmployeeExportColumn('family_composition', 'Composición familiar', 'choice'),
    EmployeeExportColumn('dependents_count', 'Personas a cargo', 'integer'),
    EmployeeExportColumn('socioeconomic_stratum', 'Estrato socioeconómico', 'choice'),
]

COLUMNS_BY_KEY = {column.key: column for column in COLUMNS}

GROUP_OF = {key: group for group, keys in EMPLOYEE_FIELD_GROUPS.items() for key in keys}

LOGO_PATH = Path(__file__).resolve().parent.parent / 'assets' / 'forjas-logo.png'

BRAND_BLUE = '#003064'

# a spreadsheet would run these as formulas; a leading quote keeps them as text
FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


# main class
class EmployeeExportService:
    FORMATS = {
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'csv': 'text/csv; charset=utf-8',
        'pdf': 'application/pdf',
    }

    @staticmethod
    def available_columns(readable: frozenset[str]) -> list[dict]:
        return [
            {
                'key': column.key,
                'label': column.label,
                'group': GROUP_OF[column.key],
                'group_label': EmployeeFieldGroup(GROUP_OF[column.key]).label,
            }
            for column in COLUMNS
            if column.key in readable and column.key in GROUP_OF
        ]

    # keeps the requested order and drops what the profile cannot read, like the filters do
    @staticmethod
    def columns_for(readable: frozenset[str], requested: list[str]) -> list[EmployeeExportColumn]:
        columns = []

        for key in dict.fromkeys(requested):
            if key in COLUMNS_BY_KEY and key in readable:
                columns.append(COLUMNS_BY_KEY[key])

        return columns

    @staticmethod
    def prepare(queryset: QuerySet) -> QuerySet:
        return queryset.select_related(
            'division', 'section', 'position', 'previous_position', 'immediate_boss', 'cost_center'
        ).prefetch_related('extensions')

    @staticmethod
    def value(employee, column: EmployeeExportColumn):
        match column.kind:
            case 'choice':
                return getattr(employee, f'get_{column.key}_display')() or None
            case 'boolean':
                value = getattr(employee, column.key)
                return None if value is None else ('Sí' if value else 'No')
            case 'relation':
                related = employee
                for part in column.source.split('.'):
                    related = getattr(related, part, None) if related is not None else None
                return related
            case 'seniority':
                seniority = employee.seniority
                if seniority is None:
                    return None
                return f'{seniority["years"]} años, {seniority["months"]} meses'
            case 'extensions':
                dates = sorted(extension.extension_date for extension in employee.extensions.all())
                return ', '.join(f'{day:%d/%m/%Y}' for day in dates) or None
            case _:
                value = getattr(employee, column.key)
                return None if value == '' else value

    @staticmethod
    def rows(queryset: QuerySet, columns: list[EmployeeExportColumn]) -> list[list]:
        return [
            [EmployeeExportService.value(employee, column) for column in columns]
            for employee in queryset
        ]

    @staticmethod
    def display(value, kind: str) -> str:
        if value is None:
            return ''

        if isinstance(value, date):
            return f'{value:%d/%m/%Y}'

        if kind == 'money':
            return '$ ' + f'{Decimal(value):,.0f}'.replace(',', '.')

        return str(value)

    @staticmethod
    def file_name(file_format: str) -> str:
        return f'empleados-{timezone.localdate():%Y-%m-%d}.{file_format}'

    @staticmethod
    def render(
        file_format: str,
        title: str,
        description: str,
        columns: list[EmployeeExportColumn],
        rows: list[list],
    ) -> bytes:
        match file_format:
            case 'xlsx':
                return EmployeeExportService._to_xlsx(title, description, columns, rows)
            case 'csv':
                return EmployeeExportService._to_csv(columns, rows)
            case _:
                return EmployeeExportService._to_pdf(title, description, columns, rows)

    @staticmethod
    def _safe_text(value):
        if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
            return f"'{value}"

        return value

    @staticmethod
    def _to_xlsx(title, description, columns, rows) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Empleados'
        sheet.append([column.label for column in columns])

        header_fill = PatternFill('solid', fgColor=BRAND_BLUE[1:])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill

        for row in rows:
            sheet.append([EmployeeExportService._safe_text(value) for value in row])

        for index, column in enumerate(columns, start=1):
            column_letter = get_column_letter(index)

            if column.kind == 'date':
                for cell in sheet[column_letter][1:]:
                    cell.number_format = 'DD/MM/YYYY'
            elif column.kind == 'money':
                for cell in sheet[column_letter][1:]:
                    cell.number_format = '"$" #,##0'

            longest = max(
                [len(column.label)] + [len(EmployeeExportService.display(row[index - 1], column.kind))
                                       for row in rows]
            )
            sheet.column_dimensions[column_letter].width = min(max(longest + 2, 10), 60)

        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions

        about = workbook.create_sheet('Reporte')
        about.append(['Reporte', title])
        about.append(['Filtros', description or 'Sin filtros'])
        about.append(['Generado', f'{timezone.localtime():%d/%m/%Y %H:%M}'])
        about.append(['Filas', len(rows)])
        about.column_dimensions['A'].width = 12
        about.column_dimensions['B'].width = 100

        buffer = io.BytesIO()
        workbook.save(buffer)

        return buffer.getvalue()

    # semicolons and a BOM, which is how Excel opens a CSV in Spanish without mangling accents
    @staticmethod
    def _to_csv(columns, rows) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=';')
        writer.writerow([column.label for column in columns])

        for row in rows:
            writer.writerow([
                EmployeeExportService._safe_text(
                    EmployeeExportService._csv_value(value, column.kind)
                )
                for value, column in zip(row, columns)
            ])

        return ('﻿' + buffer.getvalue()).encode('utf-8')

    @staticmethod
    def _csv_value(value, kind: str):
        if kind == 'money' and value is not None:
            return f'{Decimal(value):.2f}'.replace('.', ',')

        return EmployeeExportService.display(value, kind)

    @staticmethod
    def _escape(text: str) -> str:
        return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    @staticmethod
    def _to_pdf(title, description, columns, rows) -> bytes:
        buffer = io.BytesIO()
        page_width, page_height = landscape(letter)
        margin = 1.2 * cm
        generated = timezone.localtime()

        def draw_frame(canvas, document):
            canvas.saveState()
            canvas.setFillColor(colors.HexColor(BRAND_BLUE))
            canvas.rect(0, page_height - 1.6 * cm, page_width, 1.6 * cm, stroke=0, fill=1)
            canvas.drawImage(
                str(LOGO_PATH), margin, page_height - 1.25 * cm, width=3.2 * cm, height=0.95 * cm,
                mask='auto', preserveAspectRatio=True,
            )
            canvas.setFillColor(colors.white)
            canvas.setFont('Helvetica-Bold', 11)
            canvas.drawRightString(page_width - margin, page_height - 0.95 * cm, 'SIGTH')
            canvas.setFillColor(colors.HexColor('#6c757d'))
            canvas.setFont('Helvetica', 7.5)
            canvas.drawString(margin, 0.7 * cm, f'Generado el {generated:%d/%m/%Y %H:%M}')
            canvas.drawRightString(page_width - margin, 0.7 * cm, f'Página {document.page}')
            canvas.restoreState()

        document = SimpleDocTemplate(
            buffer,
            pagesize=(page_width, page_height),
            leftMargin=margin,
            rightMargin=margin,
            topMargin=2.2 * cm,
            bottomMargin=1.4 * cm,
            title=title,
        )

        font_size = 7 if len(columns) <= 10 else 6 if len(columns) <= 16 else 5
        cell_style = ParagraphStyle('cell', fontName='Helvetica', fontSize=font_size,
                                    leading=font_size + 1.5)
        head_style = ParagraphStyle('head', parent=cell_style, fontName='Helvetica-Bold',
                                    textColor=colors.white)

        data = [[Paragraph(column.label, head_style) for column in columns]]
        data += [
            [
                Paragraph(EmployeeExportService._escape(EmployeeExportService.display(value, column.kind)), cell_style)
                for value, column in zip(row, columns)
            ]
            for row in rows
        ]

        table = LongTable(data, repeatRows=1, colWidths=[document.width / len(columns)] * len(columns))
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(BRAND_BLUE)),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6f7f9')]),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#dee2e6')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))

        title_style = ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=13, leading=16,
                                     textColor=colors.HexColor(BRAND_BLUE))
        meta_style = ParagraphStyle('meta', fontName='Helvetica', fontSize=8, leading=10,
                                    textColor=colors.HexColor('#495057'))
        rows_label = '1 empleado' if len(rows) == 1 else f'{len(rows)} empleados'

        document.build(
            [
                Paragraph(EmployeeExportService._escape(title), title_style),
                Paragraph(EmployeeExportService._escape(f'Filtros: {description or "sin filtros"} · {rows_label}'),
                          meta_style),
                Spacer(0, 0.3 * cm),
                table,
            ],
            onFirstPage=draw_frame,
            onLaterPages=draw_frame,
        )

        return buffer.getvalue()
